from __future__ import annotations

import csv
import hashlib
import json
import re
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import httpx
from bs4 import BeautifulSoup
from docx import Document
from lxml import etree
from openpyxl import load_workbook
from pypdf import PdfReader

from app.config import Settings
from app.schemas import GraphEdge, GraphNode
from app.services.snapshot import SnapshotContext, SnapshotManager
from app.storage.graph import KuzuGraphStore

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".csv",
    ".xlsx",
}


@dataclass
class ParsedDocument:
    text: str
    sections: list[tuple[str, str]]
    warnings: list[str]


class DocumentParser:
    def parse(self, path: Path) -> ParsedDocument:
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"unsupported file type: {suffix}")
        if suffix == ".pdf":
            return self._parse_pdf(path)
        if suffix == ".docx":
            return self._parse_docx(path)
        if suffix in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            return ParsedDocument(text, [("文本", text)], [])
        if suffix == ".json":
            payload = json.loads(path.read_text(encoding="utf-8"))
            text = json.dumps(payload, ensure_ascii=False, indent=2)
            return ParsedDocument(text, [("JSON", text)], [])
        if suffix == ".xml":
            tree = etree.parse(str(path))
            text = "\n".join(item.strip() for item in tree.getroot().itertext() if item.strip())
            return ParsedDocument(text, [("XML", text)], [])
        if suffix in {".html", ".htm"}:
            return self._parse_html(path)
        if suffix == ".csv":
            return self._parse_csv(path)
        if suffix == ".xlsx":
            return self._parse_xlsx(path)
        raise ValueError(f"unsupported file type: {suffix}")

    @staticmethod
    def _parse_pdf(path: Path) -> ParsedDocument:
        reader = PdfReader(str(path))
        sections: list[tuple[str, str]] = []
        warnings: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append((f"第 {index} 页", text))
            else:
                warnings.append(f"第 {index} 页未提取到文本，可能是扫描件，首版不执行 OCR。")
        return ParsedDocument("\n\n".join(text for _, text in sections), sections, warnings)

    @staticmethod
    def _parse_docx(path: Path) -> ParsedDocument:
        document = Document(str(path))
        sections: list[tuple[str, str]] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if text:
                sections.append((f"段落 {index}", text))
        for table_index, table in enumerate(document.tables, start=1):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            table_text = "\n".join(rows).strip()
            if table_text:
                sections.append((f"表格 {table_index}", table_text))
        return ParsedDocument("\n\n".join(text for _, text in sections), sections, [])

    @staticmethod
    def _parse_html(path: Path) -> ParsedDocument:
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
        sections: list[tuple[str, str]] = []
        for table_index, table in enumerate(soup.find_all("table"), start=1):
            rows = []
            for row in table.find_all("tr"):
                cells = [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                sections.append((f"HTML 表格 {table_index}", "\n".join(rows)))
        body_text = soup.get_text("\n", strip=True)
        if body_text:
            sections.append(("正文", body_text))
        return ParsedDocument("\n\n".join(text for _, text in sections), sections, [])

    @staticmethod
    def _parse_csv(path: Path) -> ParsedDocument:
        with path.open(encoding="utf-8", errors="replace", newline="") as handle:
            rows = list(csv.reader(handle))
        text = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
        return ParsedDocument(text, [("CSV 表格", text)], [])

    @staticmethod
    def _parse_xlsx(path: Path) -> ParsedDocument:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sections: list[tuple[str, str]] = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    rows.append(" | ".join(values))
            if rows:
                sections.append((f"工作表 {sheet.title}", "\n".join(rows)))
        return ParsedDocument("\n\n".join(text for _, text in sections), sections, [])


def chunk_sections(
    sections: list[tuple[str, str]],
    max_chars: int = 900,
    overlap: int = 120,
) -> list[tuple[str, str]]:
    output: list[tuple[str, str]] = []
    for locator, text in sections:
        paragraphs = [item.strip() for item in re.split(r"\n{2,}|\n", text) if item.strip()]
        buffer = ""
        chunk_index = 1
        for paragraph in paragraphs:
            if len(buffer) + len(paragraph) + 1 <= max_chars:
                buffer = f"{buffer}\n{paragraph}".strip()
                continue
            if buffer:
                output.append((f"{locator} / 片段 {chunk_index}", buffer))
                chunk_index += 1
                buffer = buffer[-overlap:] + "\n" + paragraph
            else:
                start = 0
                while start < len(paragraph):
                    output.append(
                        (f"{locator} / 片段 {chunk_index}", paragraph[start : start + max_chars])
                    )
                    chunk_index += 1
                    start += max_chars - overlap
                buffer = ""
        if buffer:
            output.append((f"{locator} / 片段 {chunk_index}", buffer))
    return output


class CandidateExtractor:
    def __init__(self, graph: KuzuGraphStore):
        self.graph = graph
        self.nodes = graph.all_nodes()

    def extract(self, chunk_id: str, text: str, use_llm: bool, settings: Settings) -> list[dict[str, Any]]:
        if use_llm and settings.llm_api_key:
            return self._extract_with_llm(chunk_id, text, settings)
        return self._extract_with_rules(chunk_id, text)

    def _extract_with_rules(self, chunk_id: str, text: str) -> list[dict[str, Any]]:
        matches: list[tuple[int, GraphNode]] = []
        lowered = text.lower()
        for node in self.nodes:
            aliases = [node.name, *node.aliases]
            for alias in aliases:
                position = lowered.find(alias.lower())
                if position >= 0:
                    matches.append((position, node))
                    break
        matches.sort(key=lambda item: item[0])
        entities = []
        for _, node in matches:
            if node not in entities:
                entities.append(node)
        if len(entities) < 2:
            return []

        candidates: list[dict[str, Any]] = []
        relation_patterns = [
            ("含有", "CONTAINS"),
            ("参与", "PARTICIPATES_IN"),
            ("影响", "POTENTIALLY_AFFECTS"),
            ("导致", "CAN_CAUSE"),
            ("相互作用", "INTERACTS_WITH"),
            ("用于评估", "ASSESSES"),
        ]
        for relation_index in range(min(6, len(entities) - 1)):
            left = entities[relation_index]
            right = entities[relation_index + 1]
            predicate = "RELATED_TO"
            for phrase, candidate_predicate in relation_patterns:
                if phrase in text:
                    predicate = candidate_predicate
                    break
            candidates.append(
                {
                    "id": str(uuid.uuid4()),
                    "candidate_type": "edge",
                    "chunk_id": chunk_id,
                    "source_id": left.id,
                    "target_id": right.id,
                    "predicate": predicate,
                    "confidence": 0.45,
                    "extractor": "rules",
                    "status": "pending",
                }
            )
        return candidates

    def _extract_with_llm(
        self,
        chunk_id: str,
        text: str,
        settings: Settings,
    ) -> list[dict[str, Any]]:
        entity_catalog = [
            {"id": node.id, "name": node.name, "aliases": node.aliases}
            for node in self.nodes[:200]
        ]
        prompt = (
            "你只做离线结构化抽取，不回答问题。"
            "从文本中找出实体和关系，只能使用给定实体ID。"
            "返回JSON数组，每项包含source_id,target_id,predicate,confidence。\n"
            f"实体目录：{json.dumps(entity_catalog, ensure_ascii=False)}\n"
            f"文本：{text}"
        )
        response = httpx.post(
            f"{settings.llm_base_url.rstrip('/')}/chat/completions",
            headers={
                "Authorization": f"Bearer {settings.llm_api_key}",
                "Content-Type": "application/json",
            },
            json={
                "model": settings.llm_model,
                "temperature": 0,
                "messages": [
                    {
                        "role": "system",
                        "content": "只输出JSON，不提供医学建议。",
                    },
                    {"role": "user", "content": prompt},
                ],
            },
            timeout=60,
        )
        response.raise_for_status()
        content = response.json()["choices"][0]["message"]["content"]
        content = content.strip().removeprefix("```json").removesuffix("```").strip()
        parsed = json.loads(content)
        candidates = []
        for item in parsed if isinstance(parsed, list) else []:
            if item.get("source_id") not in {node.id for node in self.nodes}:
                continue
            if item.get("target_id") not in {node.id for node in self.nodes}:
                continue
            candidates.append(
                {
                    "id": str(uuid.uuid4()),
                    "candidate_type": "edge",
                    "chunk_id": chunk_id,
                    "source_id": item["source_id"],
                    "target_id": item["target_id"],
                    "predicate": str(item.get("predicate", "RELATED_TO")).upper(),
                    "confidence": float(item.get("confidence", 0.5)),
                    "extractor": "llm",
                    "status": "pending",
                }
            )
        return candidates


class IngestionService:
    def __init__(self, settings: Settings, manager: SnapshotManager):
        self.settings = settings
        self.manager = manager
        self.parser = DocumentParser()

    def ingest(
        self,
        path: Path,
        source_title: str | None = None,
        source_type: str = "document",
        use_llm: bool = False,
        activate: bool = False,
    ) -> dict[str, Any]:
        if not path.exists() or not path.is_file():
            raise FileNotFoundError(f"input file not found: {path}")
        current = self.manager.bootstrap()
        base_version = current.version
        current.close()
        snapshot = self.manager.create_derived_snapshot(base_version)
        warnings: list[str] = []
        try:
            parsed = self.parser.parse(path)
            warnings.extend(parsed.warnings)
            chunks = chunk_sections(parsed.sections)
            if not chunks:
                raise ValueError("document produced no text chunks")
            source_id = f"SRC-{uuid.uuid4().hex[:12].upper()}"
            document_id = f"DOC-{uuid.uuid4().hex[:12].upper()}"
            title = source_title or path.stem
            checksum = hashlib.sha256(path.read_bytes()).hexdigest()
            snapshot.metadata.add_source(
                source_id,
                title,
                source_type,
                publisher="本地导入",
                published_at="",
                url="",
                is_demo=False,
                metadata={"imported_path": str(path)},
            )
            snapshot.metadata.add_document(
                document_id,
                source_id,
                title,
                path.suffix.lower(),
                checksum,
                {"filename": path.name},
            )
            extractor = CandidateExtractor(snapshot.graph)
            candidates: list[dict[str, Any]] = []
            job_id = f"JOB-{uuid.uuid4().hex[:12].upper()}"
            for index, (locator, text) in enumerate(chunks, start=1):
                chunk_id = f"{source_id.lower()}-chunk-{index:04d}"
                evidence_id = f"{source_id.lower()}-evidence-{index:04d}"
                snapshot.metadata.add_chunk(
                    chunk_id,
                    document_id,
                    source_id,
                    text,
                    locator,
                    [],
                    {"import_job": job_id},
                )
                snapshot.metadata.add_evidence(
                    evidence_id,
                    chunk_id,
                    source_id,
                    text[:500],
                    "imported-unreviewed",
                    0.5,
                    {"import_job": job_id},
                )
                snapshot.vectors.upsert(
                    chunk_id,
                    text,
                    {"source_id": source_id, "is_demo": False, "job_id": job_id},
                )
                candidates.extend(extractor.extract(chunk_id, text, use_llm, self.settings))

            snapshot.metadata.add_review_candidates(job_id, candidates)
            manifest_path = snapshot.root / "manifest.json"
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            counts = snapshot.metadata.counts()
            manifest.update(
                {
                    "source_count": counts["sources"],
                    "chunk_count": counts["chunks"],
                    "entity_count": snapshot.graph.stats()["nodes"],
                    "relation_count": snapshot.graph.stats()["edges"],
                }
            )
            manifest_path.write_text(
                json.dumps(manifest, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            if activate:
                snapshot.close()
                self.manager.activate(snapshot.version)
            return {
                "job_id": job_id,
                "snapshot_version": snapshot.version,
                "documents": 1,
                "chunks": len(chunks),
                "candidates": len(candidates),
                "warnings": warnings,
            }
        finally:
            snapshot.close()

    @staticmethod
    def apply_review(
        snapshot: SnapshotContext,
        candidate_ids: list[str],
        decision: str,
        reviewer: str = "local",
    ) -> int:
        items = snapshot.metadata.list_review_items("pending")
        selected = [
            item
            for item in items
            if not candidate_ids or item["id"] in candidate_ids
        ]
        if decision == "approve":
            for item in selected:
                payload = item["payload"]
                if payload.get("candidate_type") != "edge":
                    continue
                edge = GraphEdge(
                    id=f"reviewed-{payload['id'].replace('-', '')[:16]}",
                    source=payload["source_id"],
                    target=payload["target_id"],
                    predicate=payload["predicate"],
                    evidence_ids=[payload["chunk_id"].replace("chunk", "evidence")],
                    evidence_level="reviewed-import",
                    confidence=float(payload.get("confidence", 0.5)),
                    is_demo=False,
                )
                snapshot.graph.add_edge(edge)
        return snapshot.metadata.decide_review(candidate_ids, decision, reviewer)

